from pathlib import Path
from io import BytesIO

from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document

# This is your PDF/TXT/Excel extraction layer.

MAX_FILE_SIZE = 3 * 1024 * 1024  # 3 MB


def extract_pdf_text(file_bytes):
    """
    Extract text from a PDF while preserving page information.
    """

    reader = PdfReader(BytesIO(file_bytes))

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):

        text = page.extract_text() or ""

        pages.append({
            "page": page_number,
            "text": text
        })

    return pages


def extract_txt_text(file_bytes):
    """
    Extract text from a TXT file.
    """

    text = file_bytes.decode(
        "utf-8",
        errors="ignore"
    )

    return [{
        "page": 1,
        "text": text
    }]

def extract_docx_text(file_bytes):
    """
    Extract text from a DOCX file.

    Paragraphs are combined into a single logical page
    because DOCX does not have a reliable page structure
    at this extraction layer.
    """

    document = Document(
        BytesIO(file_bytes)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also extract text from tables
    for table in document.tables:

        for row in table.rows:

            values = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    values.append(text)

            if values:

                paragraphs.append(
                    " | ".join(values)
                )

    return [{
        "page": 1,
        "text": "\n".join(paragraphs)
    }]


def extract_excel_text(file_bytes):
    """
    Extract Excel workbook content.

    Each worksheet is treated as a logical page
    and its rows are converted into searchable text.
    """

    workbook = load_workbook(
        filename=BytesIO(file_bytes),
        read_only=True,
        data_only=True
    )

    pages = []

    for sheet in workbook.worksheets:

        rows = []

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if values:
                rows.append(
                    " | ".join(values)
                )

        if rows:

            pages.append({
                "page": sheet.title,
                "text": "\n".join(rows)
            })

    return pages


def validate_file_size(file_bytes):
    """
    Validate uploaded file size.
    """

    size = len(file_bytes)

    if size == 0:
        raise ValueError(
            "The uploaded file is empty."
        )

    if size > MAX_FILE_SIZE:
        raise ValueError(
            "File exceeds the maximum allowed size "
            "of 3 MB."
        )   

def extract_document(file_bytes, filename):

    validate_file_size(file_bytes)

    extension = Path(
        filename
    ).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".txt":
        return extract_txt_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    if extension in [".xlsx", ".xlsm"]:
        return extract_excel_text(file_bytes)

    raise ValueError(
        "Unsupported file type. "
        "Please upload PDF, DOCX, XLSX, or TXT."
    )


def combine_document_text(pages):
    """
    Combine extracted pages/sheets into one text string.
    """

    parts = []

    for item in pages:

        parts.append(
            f"Page/Sheet: {item['page']}\n"
            f"{item['text']}"
        )

    return "\n\n".join(parts)
